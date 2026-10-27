#!/usr/bin/env python3
"""Comfort Fresh CDP jobs. Run from the cf-cdp/ folder.

  python3 cdp_cli.py sample-data                       # write placeholder inputs
  python3 cdp_cli.py feed --as-of 2026-10-27           # build dashboard_feed.json
  python3 cdp_cli.py render                            # inline feed into ui/revenue_profit.html
  python3 cdp_cli.py build-revenue --shopify-csv orders_export.csv
  python3 cdp_cli.py pull-meta --dump ads_manager_daily.csv
  python3 cdp_cli.py build-audiences --profiles warehouse/profiles/profiles.csv --wool-available 0
  python3 cdp_cli.py push-omnisend --wave A            # dry run only by default
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sys
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from cdp import config as C  # noqa: E402
from cdp import omnisend_push, pnl, sources, waves  # noqa: E402

SAMPLE = ROOT / "data" / "sample"
WH = ROOT / "warehouse"


def _read_json(p):
    return json.loads(Path(p).read_text())


def _write_json(p, obj):
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    Path(p).write_text(json.dumps(obj, indent=2) + "\n")
    print(f"wrote {p}")


def _read_csv(p):
    with open(p, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def _write_csv(p, rows, fields=None):
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    fields = fields or list(rows[0].keys())
    with open(p, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("|".join(v) if isinstance(v, list) else v) for k, v in r.items()})
    print(f"wrote {p} ({len(rows)} rows)")


# ---------------------------------------------------------------- placeholder
def placeholder_revenue(start: date, end: date, seed: int = 20261027) -> list[dict]:
    """PLACEHOLDER Shopify-shaped rows near the ~$10k/mo 2026 runway. Not actuals."""
    rng = random.Random(seed)
    packs = [(29.0, 3, 0.30), (49.0, 6, 0.55), (89.0, 12, 0.15)]  # price, pairs, weight
    rows, d, i = [], start, 0
    while d <= end:
        weekday_f = [1.05, 1.0, 0.95, 0.95, 1.0, 1.1, 1.15][d.weekday()]
        trend = 1 + 0.0015 * i
        n = max(0, round(rng.gauss(4.9 * weekday_f * trend, 1.6)))
        gross = pairs = disc = repeat = 0.0
        for _ in range(n):
            price, pr, _w = rng.choices(packs, weights=[p[2] for p in packs])[0]
            qty = 1 if rng.random() < 0.88 else 2
            gross += price * qty
            pairs += pr * qty
            if rng.random() < 0.18:
                disc += round(price * qty * 0.15, 2)
            if rng.random() < 0.21:
                repeat += 1
        refunds = round(gross * rng.uniform(0, 0.04), 2)
        rows.append({
            "date": d.isoformat(), "orders_paid": n, "gross_sales_usd": round(gross, 2),
            "discounts_usd": round(disc, 2), "refunds_usd": refunds,
            "net_sales_usd": round(gross - disc - refunds, 2), "pairs_sold": int(pairs),
            "new_customer_orders": n - int(repeat), "repeat_customer_orders": int(repeat),
            "source": "PLACEHOLDER", "as_of": end.isoformat(),
        })
        d += timedelta(days=1)
        i += 1
    return rows


SAMPLE_WAVE_SUMMARY = {
    "as_of": "2026-10-28",
    "source": "CoS snapshot 2026-10-28 (counts only)",
    "profiles": {"total": 16508, "SUBSCRIBED": 3912, "NOT_SUBSCRIBED": 10999, "UNSUBSCRIBED": 1597},
    "waves": [
        {"wave": "A", "members_all": 26, "members_sub_eligible": 26,
         "who": "1 paid order, SUB, d28–45, prefers ankle",
         "offer": "Club Starter ankle 6@$49 alone (no COMEBACK15)",
         "status": "DRAFT_HELD_FOR_SY"},
        {"wave": "B", "members_all": 43, "members_sub_eligible": 15,
         "who": "Ankle-only, d46–59", "offer": "Style switch toward crew", "status": "PLANNED"},
        {"wave": "C60", "members_all": 66, "members_sub_eligible": 35,
         "who": "Ankle-only, d60–89", "offer": "Win-back / ladder", "status": "PLANNED"},
        {"wave": "C90", "members_all": 4428, "members_sub_eligible": 997,
         "who": "Ankle-only, d90+", "offer": "Long-tail win-back; COMEBACK15 once, ends ~2026-11-26, Sy OK only",
         "status": "PLANNED"},
    ],
    "wool_c": {"status": "HOLD", "sku": C.WOOL_SKU, "sku_available": None,
               "sub_eligible_count": None, "note": "HOLD until CF22032 Shopify available > 0"},
}

# Synthetic profiles (example.invalid) that exercise every wave rule.
SAMPLE_PROFILES = [
    # email, orders, state, days, primary, styles, ankle_pref
    ("a.one@example.invalid", 1, "SUBSCRIBED", 31, "ankle", "ankle", True),
    ("a.two@example.invalid", 1, "SUBSCRIBED", 44, "ankle", "ankle", True),
    ("a.notsub@example.invalid", 1, "NOT_SUBSCRIBED", 33, "ankle", "ankle", True),  # -> none (A needs SUB)
    ("b.one@example.invalid", 2, "SUBSCRIBED", 50, "ankle", "ankle", True),
    ("b.notsub@example.invalid", 1, "NOT_SUBSCRIBED", 52, "ankle", "ankle", False),
    ("c60.one@example.invalid", 1, "SUBSCRIBED", 75, "ankle", "ankle", True),
    ("c60.unsub@example.invalid", 3, "UNSUBSCRIBED", 61, "ankle", "ankle", False),
    ("c90.one@example.invalid", 1, "SUBSCRIBED", 140, "ankle", "ankle", True),
    ("c90.notsub@example.invalid", 2, "NOT_SUBSCRIBED", 400, "ankle", "ankle", False),
    ("crew.buyer@example.invalid", 2, "SUBSCRIBED", 70, "mixed", "ankle|crew", True),  # not ankle-only
    ("fresh.buyer@example.invalid", 1, "SUBSCRIBED", 9, "crew", "crew", False),
    ("wool.owner@example.invalid", 3, "SUBSCRIBED", 95, "mixed", "ankle|crew|wool", True),
]


def cmd_sample_data(a):
    as_of = date.fromisoformat(a.as_of)
    _write_json(SAMPLE / "revenue_daily.placeholder.json",
                placeholder_revenue(date(as_of.year, as_of.month, 1) - timedelta(days=150), as_of))
    _write_json(SAMPLE / "meta_daily.json", [])  # empty on purpose: Meta not connected
    _write_json(SAMPLE / "wave_membership_summary.2026-10-28.json", SAMPLE_WAVE_SUMMARY)
    profiles = []
    for email, n, state, days, prim, styles, ap in SAMPLE_PROFILES:
        last = as_of - timedelta(days=days)
        profiles.append({
            "email": email, "shopify_customer_id": "", "first_name": "", "last_name": "",
            "email_marketing_state": state, "order_count_paid": n,
            "lifetime_revenue_usd": 49 * n, "first_order_at": "", "last_order_at": last.isoformat(),
            "days_since_last": days, "primary_style": prim, "styles_bought": styles,
            "missing_next_style": "", "day_n_band": waves.day_n_band(days), "wave": "",
            "wave_sub_eligible": state == "SUBSCRIBED", "has_active_subscription": False,
            "ankle_pref": ap, "tags": "", "as_of": as_of.isoformat(), "source": "SYNTHETIC_SAMPLE",
        })
    for p in profiles:
        p["wave"] = waves.assign_wave(p)
    _write_csv(SAMPLE / "profiles.sample.csv", profiles)
    _write_csv(SAMPLE / "wave_membership.sample.csv", waves.build_membership(profiles, as_of.isoformat()))


# ---------------------------------------------------------------- jobs
def cmd_feed(a):
    rev = _read_json(a.revenue)
    meta = _read_json(a.meta) if a.meta and Path(a.meta).exists() else []
    wv = _read_json(a.waves)
    source = rev[0].get("source", "unknown") if rev else "empty"
    feed = pnl.build_feed(rev, meta, wv, date.fromisoformat(a.as_of), source)
    qpath = Path(a.revenue).with_name("revenue_quality.json")
    if source != "PLACEHOLDER" and qpath.exists():
        q = _read_json(qpath)
        known = q["sku_lines_pack_known"] / q["sku_lines"] if q["sku_lines"] else 0
        feed["notes"] = [
            f"{q['orders_in_file']} orders in export; {q['cancelled_excluded']} cancelled excluded.",
            f"Pairs estimated from SKU pack sizes ({known:.0%} of SKU lines parsed; others count as 1 pair).",
            "Repeat = customer's 2nd+ lifetime order (lifetime count minus their later orders in the export).",
            "Gross = subtotal + discounts. Shipping charged to customers is not counted as sales.",
        ]
    _write_json(a.out, feed)


def cmd_render(a):
    feed = Path(a.feed).read_text().strip()
    html = Path(a.template).read_text()
    marker = "/*__CDP_FEED__*/null"
    if marker not in html:
        sys.exit("template is missing the feed marker")
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(html.replace(marker, feed.replace("</", "<\\/")))
    print(f"wrote {a.out}")


def cmd_build_revenue(a):
    with open(a.shopify_csv, encoding="utf-8-sig") as fh:
        header = fh.readline()
    if "processed_at_utc" in header and "total_quantity" in header:  # one row per order
        rows, quality = sources.revenue_daily_from_order_summary_csv(a.shopify_csv, a.as_of)
        _write_json(Path(a.out).with_name("revenue_quality.json"), quality)
        print(f"quality: {quality}")
    else:  # Shopify Admin line-item export
        pairs_map = _read_json(a.pairs_map) if a.pairs_map else None
        rows = sources.revenue_daily_from_shopify_csv(a.shopify_csv, pairs_map, a.as_of)
    _write_json(a.out, rows)


def cmd_pull_meta(a):
    if a.adspirer:
        sys.exit("Adspirer live pull runs from a Claude session with the Adspirer Meta connector; "
                 "write rows in meta_daily shape to warehouse/meta/daily/. See docs/JOBS.md.")
    _write_json(a.out, sources.meta_daily_from_ads_manager_csv(a.dump, a.as_of))


def cmd_magnet(a):
    from cdp import magnet
    feed, aud = magnet.build_feed(a.orders, date.fromisoformat(a.as_of))
    _write_json(a.out, feed)
    out = WH / "audiences" / "magnet"
    out.mkdir(parents=True, exist_ok=True)
    actions = aud.pop("_actions", [])
    if actions:  # one next-best action per customer (PII: git-ignored)
        _write_csv(WH / "audiences" / f"next_best_action_{a.as_of}.csv", actions)
    for seg, emails in aud.items():  # PII: stays in git-ignored warehouse/
        name = seg.replace(" · ", "_").replace(" ", "_").replace("+", "plus").replace("(", "").replace(")", "")
        _write_csv(out / f"{name}_{a.as_of}.csv", [{"email": e} for e in emails], ["email"])
    p = feed["path"]
    print(f"$100k/mo modeled in {p['hits_target']}; first-order ROAS at target CPR {p['first_order_roas_at_target_cpr']}x")


def cmd_export_emails(a):
    """Workbook of customer emails per playbook segment, for Sy-approved sends."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from cdp import magnet, playbook
    rows = magnet.load(a.orders)
    as_of = date.fromisoformat(a.as_of)
    summary, actions = playbook.build(rows, as_of)
    no_email = sum(1 for x in actions if not x["email"])
    actions = [x for x in actions if x["email"]]
    # Large feet = bought sock size 13-15 in ANY order, not only their most common size.
    big = set()
    print(f"skipped {no_email} customers with no email")
    for r in rows:
        if "13-15" in (r.get("skus") or "") or "1315" in (r.get("skus") or ""):
            big.add((r.get("customer_email") or "").strip().lower())
    cols = ["email", "first_name", "sock_size", "large_feet_13_15", "lifetime_orders", "days_since_last_order",
            "last_order", "spent_10mo_usd", "next_product", "offer", "channel", "send_plan", "marketing_consent"]
    def row(x):
        return [x["email"], x["first_name"], x["size"] or "", "YES" if x["email"] in big else "",
                x["lifetime_orders"], x["days_since"], x["last_order"], x["spend_usd"], x["product"],
                x["incentive"], x["channel"], x["send"], "CHECK in Omnisend before sending"]
    wb = Workbook()
    ov = wb.active
    ov.title = "Read me"
    ov.append(["Comfort Fresh customer email lists", f"as of {a.as_of}"])
    ov.append(["Rules", "Sy OK before any send · only SUBSCRIBED contacts get marketing · max 3 marketing emails per 30 days · no Omnisend + Gmail marketing in the same week"])
    ov.append(["Large feet", "large_feet_13_15 = YES means the customer bought sock size 13-15 at least once"])
    ov.append([])
    ov.append(["Tab", "Customers", "When", "Offer"])
    labels = {s["key"]: s for s in summary["segments"]}
    order = [s["key"] for s in summary["segments"] if s["key"] not in ("other",)]
    tabs = {"service": "Refund recovery", "vip_lapsed": "VIP gone quiet", "vip_active": "VIP active",
            "second_order": "2nd order d7-27", "wave_a": "Wave A d28-45", "household": "Household",
            "bundle_clock": "Bundle sock clock", "large_feet": "Large feet 1 order", "ladder": "1 order d46-119",
            "repeat_lapsed": "2 orders lapsed", "dormant": "Dormant 120+ (C90)", "recent": "Just ordered"}
    for key in order:
        xs = sorted((x for x in actions if x["segment"] == key), key=lambda x: (-x["lifetime_orders"], -x["spend_usd"]))
        ws = wb.create_sheet(tabs.get(key, key)[:31])
        ws.append(cols)
        for x in xs:
            ws.append(row(x))
        s = labels[key]
        ov.append([ws.title, len(xs), s["when"], s["incentive"]])
    big.discard("")
    ws = wb.create_sheet("Large feet 13-15 ALL")
    ws.append(cols + ["segment"])
    for x in sorted((x for x in actions if x["email"] in big), key=lambda x: (-x["lifetime_orders"], -x["spend_usd"])):
        ws.append(row(x) + [x["segment"]])
    ov.append([ws.title, ws.max_row - 1, "Any time (pick by segment)", "Customers who bought sock size 13-15"])
    ws = wb.create_sheet("All customers")
    ws.append(cols + ["segment"])
    for x in sorted(actions, key=lambda x: x["segment"]):
        ws.append(row(x) + [x["segment"]])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 60)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(actions)} customers, {len(big)} large-feet (size 13-15)")


def cmd_build_audiences(a):
    profiles = _read_csv(a.profiles)
    rows = waves.build_membership(profiles, a.as_of)
    counts = {"total": len(profiles)}
    for p in profiles:
        s = p.get("email_marketing_state") or "UNKNOWN"
        counts[s] = counts.get(s, 0) + 1
    wool = waves.wool_hold_list(profiles, a.wool_available)
    out = Path(a.out_dir)
    _write_csv(out / f"wave_membership_{a.as_of}.csv", rows,
               ["email", "wave", "day_n_band", "email_marketing_state", "ankle_pref", "wave_sub_eligible", "as_of"])
    _write_json(out / f"wave_membership_summary_{a.as_of}.json", waves.summarize(rows, counts, a.as_of, wool))


def cmd_push_omnisend(a):
    membership = _read_csv(a.membership)
    for r in membership:
        r["wave_sub_eligible"] = str(r.get("wave_sub_eligible")).lower() == "true"
    events = []
    if a.events and Path(a.events).exists():
        events = [json.loads(line) for line in Path(a.events).read_text().splitlines() if line.strip()]
    plan = omnisend_push.plan_push(membership, a.wave, events, date.fromisoformat(a.as_of), a.wool_status)
    _write_json(a.out or WH / "audiences" / f"omnisend_plan_{a.wave}_{a.as_of}.json", plan)
    print(f"DRY RUN: would upsert {plan['would_upsert']} contacts; skipped {plan['skipped']}")
    if a.execute:
        omnisend_push.execute_push(plan, a.approved_by, a.approval_ref)


def main(argv=None):
    today = date.today().isoformat()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("sample-data"); s.add_argument("--as-of", default="2026-10-27"); s.set_defaults(fn=cmd_sample_data)

    s = sub.add_parser("feed")
    s.add_argument("--revenue", default=str(SAMPLE / "revenue_daily.placeholder.json"))
    s.add_argument("--meta", default=str(SAMPLE / "meta_daily.json"))
    s.add_argument("--waves", default=str(SAMPLE / "wave_membership_summary.2026-10-28.json"))
    s.add_argument("--as-of", default="2026-10-27")
    s.add_argument("--out", default=str(SAMPLE / "dashboard_feed.json"))
    s.set_defaults(fn=cmd_feed)

    s = sub.add_parser("render")
    s.add_argument("--feed", default=str(SAMPLE / "dashboard_feed.json"))
    s.add_argument("--template", default=str(ROOT / "ui" / "templates" / "revenue_profit.html"))
    s.add_argument("--out", default=str(ROOT / "ui" / "revenue_profit.html"))
    s.set_defaults(fn=cmd_render)

    s = sub.add_parser("build-revenue")
    s.add_argument("--shopify-csv", required=True)
    s.add_argument("--pairs-map", help="JSON {sku: pairs_per_unit}")
    s.add_argument("--as-of", default=today)
    s.add_argument("--out", default=str(WH / "revenue" / "revenue_daily.json"))
    s.set_defaults(fn=cmd_build_revenue)

    s = sub.add_parser("pull-meta")
    s.add_argument("--dump", help="Ads Manager CSV with Day breakdown")
    s.add_argument("--adspirer", action="store_true")
    s.add_argument("--as-of", default=today)
    s.add_argument("--out", default=str(WH / "meta" / "daily" / f"meta_daily_{today}.json"))
    s.set_defaults(fn=cmd_pull_meta)

    s = sub.add_parser("magnet", help="Money Magnet: findings, ranked moves, path to $100k")
    s.add_argument("--orders", required=True, help="one-row-per-order Shopify export")
    s.add_argument("--as-of", default=today)
    s.add_argument("--out", default=str(WH / "feeds" / "magnet_feed.json"))
    s.set_defaults(fn=cmd_magnet)

    s = sub.add_parser("export-emails", help="Excel workbook of customer emails per segment")
    s.add_argument("--orders", required=True)
    s.add_argument("--as-of", default=today)
    s.add_argument("--out", default=str(WH / "exports" / f"customer_email_lists_{today}.xlsx"))
    s.set_defaults(fn=cmd_export_emails)

    s = sub.add_parser("build-audiences")
    s.add_argument("--profiles", default=str(SAMPLE / "profiles.sample.csv"))
    s.add_argument("--as-of", default=today)
    s.add_argument("--wool-available", type=int, default=None)
    s.add_argument("--out-dir", default=str(WH / "audiences"))
    s.set_defaults(fn=cmd_build_audiences)

    s = sub.add_parser("push-omnisend")
    s.add_argument("--membership", default=str(SAMPLE / "wave_membership.sample.csv"))
    s.add_argument("--wave", required=True, choices=["A", "B", "C60", "C90", "WOOL_C"])
    s.add_argument("--events", default=str(WH / "events" / "events.jsonl"))
    s.add_argument("--as-of", default=today)
    s.add_argument("--wool-status", default="HOLD")
    s.add_argument("--out")
    s.add_argument("--execute", action="store_true", help="live push; needs Sy approval flags")
    s.add_argument("--approved-by")
    s.add_argument("--approval-ref")
    s.set_defaults(fn=cmd_push_omnisend)

    a = ap.parse_args(argv)
    a.fn(a)


if __name__ == "__main__":
    main()
