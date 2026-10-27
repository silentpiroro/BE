"""Offline source adapters.

- Shopify orders CSV export (Admin > Orders > Export) -> revenue_daily rows
- Meta Ads Manager CSV dump (daily breakdown)       -> meta_daily rows

Live connectors (Shopify Admin MCP run by CoS, Adspirer Meta) produce the same
row shapes; see docs/JOBS.md.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict

PAID_STATUSES = {"paid", "partially_refunded", "refunded", "partially_paid"}
_PAIRS_RE = re.compile(r"(\d+)\s*(?:-|\s)?\s*(?:pairs?|pack|pk)\b", re.I)


def pairs_in_title(title: str, pairs_by_sku: dict | None = None, sku: str = "") -> int:
    """Pairs per unit. Prefer an explicit SKU map; Store SKUs are not warehouse 1:1."""
    if pairs_by_sku and sku in pairs_by_sku:
        return int(pairs_by_sku[sku])
    m = _PAIRS_RE.search(title or "")
    return int(m.group(1)) if m else 1


def _f(v) -> float:
    try:
        return float(str(v).replace("$", "").replace(",", "") or 0)
    except ValueError:
        return 0.0


def revenue_daily_from_shopify_csv(path: str, pairs_by_sku: dict | None = None,
                                   as_of: str = "") -> list[dict]:
    orders: dict[str, dict] = {}
    with open(path, newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            name = row.get("Name")
            if not name:
                continue
            o = orders.get(name)
            if o is None and row.get("Financial Status"):
                o = orders[name] = {
                    "email": (row.get("Email") or "").strip().lower(),
                    "status": row["Financial Status"].strip().lower(),
                    "paid_at": (row.get("Paid at") or row.get("Created at") or "")[:10],
                    "subtotal": _f(row.get("Subtotal")),
                    "discount": _f(row.get("Discount Amount")),
                    "refunded": _f(row.get("Refunded Amount")),
                    "pairs": 0,
                }
            if o is not None:
                qty = int(_f(row.get("Lineitem quantity")))
                o["pairs"] += qty * pairs_in_title(row.get("Lineitem name", ""), pairs_by_sku,
                                                   row.get("Lineitem sku", ""))
    seen_emails: set[str] = set()
    days: dict[str, dict] = defaultdict(lambda: {
        "orders_paid": 0, "gross_sales_usd": 0.0, "discounts_usd": 0.0, "refunds_usd": 0.0,
        "net_sales_usd": 0.0, "pairs_sold": 0, "new_customer_orders": 0, "repeat_customer_orders": 0})
    for o in sorted(orders.values(), key=lambda x: x["paid_at"]):
        if o["status"] not in PAID_STATUSES or not o["paid_at"]:
            continue
        d = days[o["paid_at"]]
        d["orders_paid"] += 1
        d["gross_sales_usd"] += o["subtotal"] + o["discount"]
        d["discounts_usd"] += o["discount"]
        d["refunds_usd"] += o["refunded"]
        d["net_sales_usd"] += o["subtotal"] - o["refunded"]
        d["pairs_sold"] += o["pairs"]
        # New vs repeat is only as good as the export's date range.
        if o["email"] and o["email"] in seen_emails:
            d["repeat_customer_orders"] += 1
        else:
            d["new_customer_orders"] += 1
        seen_emails.add(o["email"])
    out = []
    for day in sorted(days):
        r = {"date": day, **days[day], "source": "shopify_orders_csv", "as_of": as_of}
        for k in ("gross_sales_usd", "discounts_usd", "refunds_usd", "net_sales_usd"):
            r[k] = round(r[k], 2)
        out.append(r)
    return out


# Ads Manager column names vary by export settings; first match wins.
META_COLS = {
    "date": ["Day", "Reporting starts", "Date"],
    "campaign_name": ["Campaign name", "Campaign Name"],
    "adset_name": ["Ad set name", "Ad Set Name"],
    "spend_usd": ["Amount spent (USD)", "Amount spent", "Spend"],
    "impressions": ["Impressions"],
    "clicks": ["Link clicks", "Clicks (all)", "Clicks"],
    "purchases": ["Purchases", "Website purchases"],
    "purchase_value_usd": ["Purchases conversion value", "Website purchases conversion value",
                           "Purchase conversion value"],
}


def _ads_rows(path: str) -> tuple[list[str], list[dict]]:
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h) for h in next(it)]
        return head, [dict(zip(head, ["" if v is None else v for v in r])) for r in it]
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return list(reader.fieldnames or []), list(reader)


def meta_period_from_ads_export(path: str) -> dict:
    """Totals for an Ads Manager export without a Day breakdown (one reporting period)."""
    head, rows = _ads_rows(path)
    starts = {str(r.get("Reporting starts"))[:10] for r in rows}
    ends = {str(r.get("Reporting ends"))[:10] for r in rows}
    spend = sum(_f(r.get("Amount spent (USD)")) for r in rows)
    purchases = sum(_f(r.get("Purchases")) for r in rows)
    value = sum(_f(r.get("Amount spent (USD)")) * _f(r.get("Purchase ROAS (return on ad spend)")) for r in rows)
    start, end = min(starts), max(ends)
    from datetime import date as _date
    days = (_date.fromisoformat(end) - _date.fromisoformat(start)).days + 1
    by_set = defaultdict(lambda: {"spend_usd": 0.0, "purchases": 0.0, "adds_to_cart": 0.0})
    for r in rows:
        b = by_set[r.get("Ad set name") or ""]
        b["spend_usd"] += _f(r.get("Amount spent (USD)"))
        b["purchases"] += _f(r.get("Purchases"))
        b["adds_to_cart"] += _f(r.get("Adds to cart"))
    return {"start": start, "end": end, "days": days, "spend_usd": round(spend, 2), "purchases": int(purchases),
            "purchase_value_usd": round(value, 2), "cpr_usd": round(spend / purchases, 2) if purchases else None,
            "purchase_roas": round(value / spend, 2) if spend else None,
            "adds_to_cart": int(sum(_f(r.get("Adds to cart")) for r in rows)),
            "ad_sets": {k: {kk: round(vv, 2) for kk, vv in v.items()} for k, v in by_set.items() if v["spend_usd"] > 0}}


def meta_daily_from_ads_manager_csv(path: str, as_of: str = "") -> list[dict]:
    out = []
    head, rows_in = _ads_rows(path)
    if True:
        cols = {k: next((c for c in opts if c in head), None) for k, opts in META_COLS.items()}
        if not cols["date"] or not cols["spend_usd"]:
            raise ValueError("Meta dump needs a daily breakdown (Day) and Amount spent column")
        for row in rows_in:
            g = lambda k: row.get(cols[k]) if cols[k] else None
            spend = _f(g("spend_usd"))
            purchases = int(_f(g("purchases")))
            value = _f(g("purchase_value_usd"))
            out.append({
                "date": str(g("date") or "")[:10],
                "campaign_name": g("campaign_name") or "",
                "adset_name": g("adset_name") or "",
                "spend_usd": round(spend, 2),
                "impressions": int(_f(g("impressions"))),
                "clicks": int(_f(g("clicks"))),
                "purchases": purchases,
                "purchase_value_usd": round(value, 2),
                "cpr_usd": round(spend / purchases, 2) if purchases else None,
                "purchase_roas": round(value / spend, 2) if spend else None,
                "source": "ads_manager_dump",
                "as_of": as_of,
            })
    return out


# ---------------------------------------------------------------- order-summary export
# One row per order (order_id, processed_at_utc, financial_status, subtotal_original,
# total_discounts_original, total_refunded, customer_id, customer_lifetime_orders,
# total_quantity, skus "a; b; c"). subtotal_original is AFTER discounts.
_PACK_PATTERNS = [
    re.compile(r"PK(\d+)", re.I),                 # CF22032-PK3-Black
    re.compile(r"(\d+)\s*PR\b", re.I),            # ADDON-BAMBOO-CREW-2PR
    re.compile(r"^CF-B[A-Z]{2}-(\d+)-", re.I),    # CF-BDA-20-BLK-1315
    re.compile(r"^[A-Za-z0-9]+-(\d+)-[A-Za-z]", re.I),  # CF17180-8-BLK
    re.compile(r"(\d+)\s*pairs?\b", re.I),        # "Bonus 6 Pairs"
]
_SIZE_FIRST = re.compile(r"^[A-Za-z0-9]+-(?:0?911|1013|1315)-((?:[A-Z]\d+)+)P?\b", re.I)
_HEAD_COUNTS = re.compile(r"^[A-Za-z0-9]+-((?:\d+[A-Za-z]+)+)")  # CF17180-4Black, -3Black3White, -4Pack


def pairs_per_unit_from_sku(sku: str) -> int | None:
    """Best-effort pack size from a Comfort Fresh SKU or title. None = unknown."""
    sku = sku.strip()
    m = _SIZE_FIRST.match(sku)                    # CF17180-1315-B12W6-AID, CF17180-911-B6P
    if m:
        n = sum(int(x) for x in re.findall(r"\d+", m.group(1)))
        return n if 0 < n <= 48 else None
    m = _HEAD_COUNTS.match(sku)
    if m:
        n = sum(int(n) for n in re.findall(r"(\d+)[A-Za-z]+", m.group(1)))
        return n if 0 < n <= 48 else None
    for pat in _PACK_PATTERNS:
        m = pat.search(sku)
        if m:
            n = int(m.group(1))
            return n if 0 < n <= 48 else None
    return None


def _local_date(ts: str, tz: str) -> str:
    from datetime import datetime
    from zoneinfo import ZoneInfo
    dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
    return dt.astimezone(ZoneInfo(tz)).date().isoformat()


def revenue_daily_from_order_summary_csv(path: str, as_of: str = "", tz: str = "America/Chicago",
                                         default_pack: int = 1) -> tuple[list[dict], dict]:
    """Returns (revenue_daily rows, quality report). Cancelled orders are excluded."""
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    # Customer's order number at the time = lifetime orders now - their later orders in this file.
    by_cust: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        if r.get("customer_id"):
            by_cust[r["customer_id"]].append(r)
    seq: dict[str, int] = {}
    for cust, orders in by_cust.items():
        orders.sort(key=lambda r: r["processed_at_utc"])
        life = int(_f(orders[-1].get("customer_lifetime_orders")) or len(orders))
        for i, r in enumerate(orders):
            seq[r["order_id"]] = max(life - (len(orders) - 1 - i), i + 1)
    days: dict[str, dict] = defaultdict(lambda: {
        "orders_paid": 0, "gross_sales_usd": 0.0, "discounts_usd": 0.0, "refunds_usd": 0.0,
        "net_sales_usd": 0.0, "pairs_sold": 0, "new_customer_orders": 0, "repeat_customer_orders": 0})
    q = {"orders_in_file": len(rows), "cancelled_excluded": 0, "not_paid_excluded": 0,
         "sku_lines": 0, "sku_lines_pack_known": 0}
    for r in rows:
        if r.get("cancelled_at"):
            q["cancelled_excluded"] += 1
            continue
        if r["financial_status"].strip().lower() not in PAID_STATUSES:
            q["not_paid_excluded"] += 1
            continue
        d = days[_local_date(r["processed_at_utc"], tz)]
        sub, disc, ref = _f(r["subtotal_original"]), _f(r["total_discounts_original"]), _f(r["total_refunded"])
        d["orders_paid"] += 1
        d["gross_sales_usd"] += sub + disc
        d["discounts_usd"] += disc
        d["refunds_usd"] += min(ref, sub)
        d["net_sales_usd"] += sub - min(ref, sub)
        packs = []
        for s in (r.get("skus") or "").split(";"):
            if not s.strip() or s.strip().lower().startswith("order split"):
                continue
            q["sku_lines"] += 1
            p = pairs_per_unit_from_sku(s)
            if p is not None:
                q["sku_lines_pack_known"] += 1
            packs.append(p if p is not None else default_pack)
        units = int(_f(r.get("total_quantity")))
        avg_pack = sum(packs) / len(packs) if packs else default_pack
        d["pairs_sold"] += round(units * avg_pack)
        if seq.get(r["order_id"], 1) > 1:
            d["repeat_customer_orders"] += 1
        else:
            d["new_customer_orders"] += 1
    out = []
    for day in sorted(days):
        rr = {"date": day, **days[day], "source": "shopify_order_export", "as_of": as_of}
        for k in ("gross_sales_usd", "discounts_usd", "refunds_usd", "net_sales_usd"):
            rr[k] = round(rr[k], 2)
        out.append(rr)
    return out, q
